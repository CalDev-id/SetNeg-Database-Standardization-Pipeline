import asyncio
import json
import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.routes import data_dictionary as dictionary_routes
from api.routes import audit_db as audit_routes
from api.routes import normalization as routes
from features.audit_db import pipeline as audit_pipeline
from features.audit_db.audit_output import COMBINED_COLUMNS, buat_excel_gabungan
from features.audit_db.extract import parse_schema_ddl
from features.normalization.analysis import (
    FINDING_COLUMNS, NormalizationError, analyze_normalization, parse_response,
)
from features.normalization import pipeline
from features.normalization.output import create_normalization_excel
from shared.llm.gemini_runtime import GeminiRuntime
from shared.llm.telkom_gpt_runtime import TelkomRuntimeError
from main import app


def ddl(column_counts):
    return "\n".join(
        f"CREATE TABLE `table_{index}` (" + ", ".join(
            f"`field_{number}` INT" for number in range(count)
        ) + ");"
        for index, count in enumerate(column_counts)
    ).encode()


def finding(table="pegawai", column="jabatan_akhir"):
    return {
        "No": 1, "Nama Database": "-", "Nama Tabel": table,
        "Tahap Normalisasi": "3NF", "Objek Temuan": "3NF - Duplikasi Lintas Tabel",
        "Kolom Terdampak": f"`{column}`",
        "Uraian Ketidaksesuaian": "Atribut jabatan tersimpan ulang.",
        "Solusi Normalisasi": "Gunakan relasi ke tabel jabatan.",
        "Tabel Baru Usulan": "-", "Status": "To Do", "Catatan": "-",
    }


class FakeRuntime:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def generate_response(self, system, user):
        self.calls.append((system, user))
        return self.response(self.calls) if callable(self.response) else self.response


class NormalizationAnalysisTests(unittest.IsolatedAsyncioTestCase):
    async def test_full_schema_sent_in_one_request(self):
        parsed = parse_schema_ddl(ddl([10] * 6 + [51]))
        runtime = FakeRuntime("[]")
        progress = []
        self.assertEqual(await analyze_normalization(parsed, runtime, lambda: progress.append(1)), [])
        self.assertEqual(len(runtime.calls), 1)
        self.assertEqual(len(progress), 1)
        payload = json.loads(runtime.calls[0][1].split("Data DDL:\n", 1)[1])
        self.assertEqual(len(payload["tabel"]), 7)
        self.assertEqual(len(payload["tabel"][-1]["definisi_kolom"]), 51)
        self.assertNotIn("katalog_konteks_database", payload)

    async def test_cross_table_context_and_prompt_unchanged(self):
        source = b"CREATE TABLE `pegawai` (`id` INT, `jabatan_akhir` VARCHAR(80));\n" \
                 b"CREATE TABLE `jabatan` (`id` INT, `nama` VARCHAR(80));"
        parsed = parse_schema_ddl(source)
        runtime = FakeRuntime(json.dumps([finding()]))
        result = await analyze_normalization(parsed, runtime)
        self.assertEqual(result[0]["No"], 1)
        self.assertEqual(result[0]["Nama Tabel"], "pegawai")
        self.assertIn('"nama_tabel": "jabatan"', runtime.calls[0][1])
        self.assertIn('"nama_tabel": "pegawai"', runtime.calls[0][1])
        self.assertEqual(len(runtime.calls), 1)
        from features.normalization.normalization import SYSTEM_PROMPT_NORMALISASI
        self.assertEqual(runtime.calls[0][0], SYSTEM_PROMPT_NORMALISASI)

    async def test_composite_table_name_is_split_by_real_ddl_columns_in_one_llm_call(self):
        parsed = parse_schema_ddl(
            b"CREATE TABLE `pegawai` (`nip` VARCHAR(18), `jabatan_akhir` VARCHAR(80));\n"
            b"CREATE TABLE `jabatan` (`nip` VARCHAR(18), `kode_jabatan` VARCHAR(20));"
        )
        combined = finding(table="pegawai, jabatan", column="nip")
        combined["Kolom Terdampak"] = "`nip`, `jabatan_akhir`, `kode_jabatan`"
        runtime = FakeRuntime(json.dumps([combined]))
        result = await analyze_normalization(parsed, runtime)
        self.assertEqual(len(runtime.calls), 1)
        self.assertEqual([(row["Nama Tabel"], row["Kolom Terdampak"]) for row in result], [
            ("pegawai", "`nip`, `jabatan_akhir`"),
            ("jabatan", "`nip`, `kode_jabatan`"),
        ])
        self.assertEqual([row["No"] for row in result], [1, 2])

    async def test_composite_table_name_without_columns_on_second_table_keeps_only_owner(self):
        parsed = parse_schema_ddl(
            b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));\n"
            b"CREATE TABLE `jabatan` (`kode_jabatan` VARCHAR(20));"
        )
        result = parse_response(json.dumps([finding(table="pegawai, jabatan")]), parsed)
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]["Nama Tabel"], "pegawai")

    async def test_composite_table_accepts_qualified_column_formats(self):
        parsed = parse_schema_ddl(
            b"CREATE TABLE `pegawai` (`nip` VARCHAR(18), `jabatan_akhir` VARCHAR(80));\n"
            b"CREATE TABLE `jabatan` (`nip` VARCHAR(18), `kode_jabatan` VARCHAR(20));"
        )
        cases = (
            "pegawai: `nip`, `jabatan_akhir`; jabatan: `kode_jabatan`",
            "pegawai.`nip`, `pegawai`.`jabatan_akhir`, jabatan.`kode_jabatan`",
            "`nip` (pegawai), `jabatan_akhir` (pegawai), `kode_jabatan` (jabatan)",
        )
        for columns in cases:
            with self.subTest(columns=columns):
                item = finding(table="pegawai, jabatan")
                item["Kolom Terdampak"] = columns
                runtime = FakeRuntime(json.dumps([item]))
                result = await analyze_normalization(parsed, runtime)
                self.assertEqual(len(runtime.calls), 1)
                self.assertEqual([(row["Nama Tabel"], row["Kolom Terdampak"]) for row in result], [
                    ("pegawai", "`nip`, `jabatan_akhir`"),
                    ("jabatan", "`kode_jabatan`"),
                ])

    async def test_composite_table_skips_wrong_qualified_owner(self):
        parsed = parse_schema_ddl(
            b"CREATE TABLE `pegawai` (`nip` VARCHAR(18));\n"
            b"CREATE TABLE `jabatan` (`kode_jabatan` VARCHAR(20));"
        )
        item = finding(table="pegawai, jabatan")
        item["Kolom Terdampak"] = "jabatan.`nip`"
        self.assertEqual(parse_response(json.dumps([item]), parsed), [])

    async def test_composite_table_skips_unknown_table_or_column(self):
        parsed = parse_schema_ddl(
            b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));\n"
            b"CREATE TABLE `jabatan` (`kode_jabatan` VARCHAR(20));"
        )
        for invalid in (
            finding(table="pegawai, fiktif"),
            finding(table="pegawai, jabatan", column="fiktif"),
        ):
            with self.subTest(invalid=invalid):
                self.assertEqual(parse_response(json.dumps([invalid]), parsed), [])

    async def test_invalid_json_fails_without_another_llm_call(self):
        parsed = parse_schema_ddl(ddl([1]))
        runtime = FakeRuntime(lambda calls: "not json" if len(calls) == 1 else "[]")
        with patch("features.normalization.analysis.perbaiki_json_array", new_callable=AsyncMock, return_value=None) as repair:
            with self.assertRaisesRegex(NormalizationError, "Format respons"):
                await analyze_normalization(parsed, runtime)
        repair.assert_awaited_once_with("not json")
        self.assertEqual(len(runtime.calls), 1)

    async def test_malformed_json_is_repaired_by_llm_then_validated(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        malformed = '[{"No": 1, "Nama Tabel": "pegawai",]'
        runtime = FakeRuntime(malformed)
        with patch("features.audit_db.audit.GeminiRuntime") as repair_runtime:
            repair_runtime.return_value.generate_response.return_value = json.dumps({
                "repairable": True, "findings": [finding()],
            })
            result = await analyze_normalization(parsed, runtime)
        repair_runtime.assert_called_once_with(model="gemini-3.5-flash-lite")
        self.assertIn(malformed, repair_runtime.return_value.generate_response.call_args.args[1])
        self.assertEqual(len(runtime.calls), 1)
        self.assertEqual(result[0]["Nama Tabel"], "pegawai")

    async def test_repaired_json_skips_table_outside_ddl(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        runtime = FakeRuntime("broken JSON")
        with patch("features.normalization.analysis.perbaiki_json_array", new_callable=AsyncMock,
                   return_value=[finding(table="fiktif")]):
            self.assertEqual(await analyze_normalization(parsed, runtime), [])

    async def test_repair_llm_cannot_convert_ambiguous_output_to_empty_success(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        with patch("features.audit_db.audit.GeminiRuntime") as repair_runtime:
            repair_runtime.return_value.generate_response.return_value = '{"repairable": false}'
            with self.assertRaisesRegex(NormalizationError, "Format respons"):
                await analyze_normalization(parsed, FakeRuntime("broken JSON"))

    async def test_valid_json_with_invalid_field_is_skipped_without_repair_llm(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        runtime = FakeRuntime(json.dumps([finding(table="fiktif")]))
        with patch("features.normalization.analysis.perbaiki_json_array", new_callable=AsyncMock) as repair:
            self.assertEqual(await analyze_normalization(parsed, runtime), [])
        repair.assert_not_awaited()

    async def test_model_text_around_json_uses_llm_repair(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        response = "<think>analisis selesai</think>\n" + json.dumps([finding()]) + "\nSelesai."
        runtime = FakeRuntime(response)
        with patch("features.normalization.analysis.perbaiki_json_array", new_callable=AsyncMock,
                   return_value=[finding()]) as repair:
            result = await analyze_normalization(parsed, runtime)
        repair.assert_awaited_once_with(response)
        self.assertEqual(result[0]["Nama Tabel"], "pegawai")
        self.assertEqual(len(runtime.calls), 1)

    async def test_accepts_valid_json_object_wrapping_findings_without_retry(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        for key in ("findings", "temuan", "hasil"):
            runtime = FakeRuntime(json.dumps({key: [finding()]}))
            result = await analyze_normalization(parsed, runtime)
            self.assertEqual(result[0]["Nama Tabel"], "pegawai")
            self.assertEqual(len(runtime.calls), 1)

    async def test_trailing_json_commas_use_llm_repair(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        responses = (
            json.dumps([finding()])[:-1] + ",]",
            json.dumps({"findings": [finding()]})[:-2] + ",]}",
        )
        for response in responses:
            with self.subTest(response=response):
                runtime = FakeRuntime(response)
                with patch("features.normalization.analysis.perbaiki_json_array", new_callable=AsyncMock,
                           return_value=[finding()]) as repair:
                    result = await analyze_normalization(parsed, runtime)
                repair.assert_awaited_once_with(response)
                self.assertEqual(result[0]["Nama Tabel"], "pegawai")
                self.assertEqual(len(runtime.calls), 1)

    async def test_extra_fields_and_missing_number_do_not_fail_normalization(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        item = finding()
        del item["No"]
        item["Penjelasan Tambahan"] = "Tidak digunakan oleh laporan."
        result = await analyze_normalization(parsed, FakeRuntime(json.dumps([item])))
        self.assertEqual(result[0]["No"], 1)
        self.assertNotIn("Penjelasan Tambahan", result[0])

    async def test_malformed_response_explains_failure_without_partial_success(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`jabatan_akhir` VARCHAR(80));")
        cases = (
            ('[{"Nama Tabel": "pegawai"}', "Sintaks JSON salah"),
            ('[{Nama Tabel: "pegawai"}]', "Sintaks JSON salah"),
            ('{"message": "tidak ada temuan"}', "tidak memiliki daftar temuan"),
        )
        for response, reason in cases:
            with self.subTest(response=response):
                runtime = FakeRuntime(response)
                with patch("features.normalization.analysis.perbaiki_json_array", new_callable=AsyncMock,
                           return_value=None):
                    with self.assertRaisesRegex(NormalizationError, reason):
                        await analyze_normalization(parsed, runtime)
                self.assertEqual(len(runtime.calls), 1)

    async def test_skips_unknown_table_column_and_still_reports_llm_failure(self):
        parsed = parse_schema_ddl(b"CREATE TABLE `pegawai` (`id` INT, `jabatan_akhir` VARCHAR(80));")
        for bad in (finding(table="fiktif"), finding(column="fiktif")):
            self.assertEqual(parse_response(json.dumps([bad]), parsed), [])
        with self.assertRaises(NormalizationError):
            await analyze_normalization(parsed, FakeRuntime(""))
        def fail(_calls):
            raise TelkomRuntimeError("layanan tidak tersedia")
        with self.assertRaisesRegex(NormalizationError, "LLM gagal"):
            await analyze_normalization(parsed, FakeRuntime(fail))


class NormalizationRouteTests(unittest.IsolatedAsyncioTestCase):
    async def test_homepage_is_not_cached_between_deployments(self):
        with TestClient(app) as client:
            response = client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertIn("/static/styles.css?v=41", response.text)
        self.assertIn("/static/app.js?v=53", response.text)

    async def test_combined_excel_colors_normalization_blue_and_audit_by_existing_rules(self):
        base = {key: "-" for _, key in COMBINED_COLUMNS}
        findings = [
            {**base, "no": 1, "jenis_temuan": "Audit Playbook", "solusi": "Hapus kolom lama"},
            {**base, "no": 2, "jenis_temuan": "Audit Playbook", "objek_temuan": "Tabel metadata database"},
            {**base, "no": 3, "jenis_temuan": "Audit Playbook", "objek_temuan": "Foreign Key"},
            {**base, "no": 4, "jenis_temuan": "Normalisasi", "solusi": "Hapus kolom duplikat"},
        ]
        sheet = load_workbook(BytesIO(buat_excel_gabungan(findings))).active
        self.assertEqual(sheet.max_column, 10)
        for row_number, expected in ((2, "FFFFE2DF"), (3, "FFFFF1BF"), (4, None), (5, "FFEAF4FF")):
            for cell in sheet[row_number]:
                with self.subTest(row=row_number, column=cell.column):
                    self.assertEqual(cell.fill.fgColor.rgb if cell.fill.patternType else None, expected)

    async def test_combined_report_uses_one_sheet_and_ten_columns(self):
        audit = SimpleNamespace(status="completed", normalization_job_id="normal", audit_result={
            "summary": {"table_count": 1, "column_count": 2, "tables": ["pegawai"]},
            "findings": [{
                "no": 1, "nama_database": "-", "nama_tabel": "pegawai",
                "objek_temuan": "Audit", "aturan_playbook": "2a", "uraian_ketidaksesuaian": "Salah",
                "solusi": "Perbaiki", "status": "To Do", "catatan": "-",
            }],
        })
        normal_finding = finding()
        normal_finding["Tabel Baru Usulan"] = "master_jabatan (id, nama)"
        normalization = SimpleNamespace(audit_job_id="audit", status="completed", findings=[normal_finding])
        with patch.object(audit_routes, "find_spreadsheet_job", return_value=audit), \
             patch.object(audit_routes, "find_normalization_job", return_value=normalization):
            result = await audit_routes.get_combined_findings("audit")
            response = await audit_routes.download_combined_findings("audit")
            with TestClient(app) as client:
                self.assertEqual(client.get("/api/v1/audit/spreadsheet/audit/combined-findings").status_code, 200)
                self.assertEqual(client.get("/api/v1/audit/spreadsheet/audit/download/combined").status_code, 200)
        self.assertEqual([row["jenis_temuan"] for row in result["findings"]], ["Audit Playbook", "Normalisasi"])
        self.assertEqual([row["no"] for row in result["findings"]], [1, 2])
        self.assertNotIn("jenis_temuan", audit.audit_result["findings"][0])
        self.assertIn("`jabatan_akhir`", result["findings"][1]["uraian_ketidaksesuaian"])
        self.assertIn("master_jabatan", result["findings"][1]["solusi"])
        workbook = load_workbook(BytesIO(response.body))
        self.assertEqual(workbook.sheetnames, ["Temuan Database"])
        sheet = workbook.active
        self.assertEqual([cell.value for cell in sheet[1]], [label for label, _ in COMBINED_COLUMNS])
        self.assertEqual(sheet.max_column, 10)
        self.assertEqual(sheet["J2"].value, "Audit Playbook")
        self.assertEqual(sheet["J3"].value, "Normalisasi")

    async def test_combined_report_waits_for_matching_normalization(self):
        audit = SimpleNamespace(status="completed", normalization_job_id="normal", audit_result={"findings": []})
        with patch.object(audit_routes, "find_spreadsheet_job", return_value=audit):
            for normalization in (
                None,
                SimpleNamespace(audit_job_id="other", status="completed", findings=[]),
                SimpleNamespace(audit_job_id="audit", status="failed", findings=None, error="Gagal"),
            ):
                with patch.object(audit_routes, "find_normalization_job", return_value=normalization):
                    with self.assertRaises(HTTPException) as caught:
                        await audit_routes.download_combined_findings("audit")
                    self.assertEqual(caught.exception.status_code, 409)

    async def test_normalization_findings_can_be_displayed_while_audit_runs(self):
        item = finding()
        job = SimpleNamespace(status="completed", findings=[item], table_count=2, column_count=3)
        with patch.object(routes, "find_normalization_job", return_value=job):
            result = await routes.get_normalization_findings("normal")
        self.assertEqual(result["findings"], [item])
        self.assertEqual(result["display_findings"][0]["jenis_temuan"], "Normalisasi")
        self.assertEqual(result["display_findings"][0]["nama_tabel"], "pegawai")
        self.assertEqual(result["display_findings"][0]["aturan_playbook"], "Normalisasi Database (3NF)")

    async def test_normalization_runtime_disables_transport_retries_and_fallbacks(self):
        runtime = SimpleNamespace(max_attempts=5)
        with patch.object(pipeline, "get_runtime", return_value=runtime):
            self.assertIs(pipeline.get_normalization_runtime("telkom"), runtime)
        self.assertEqual(runtime.max_attempts, 1)

        gemini = GeminiRuntime.__new__(GeminiRuntime)
        gemini.model = "model-utama"
        gemini.models = ("model-utama", "model-cadangan")
        with patch.object(pipeline, "get_runtime", return_value=gemini):
            pipeline.get_normalization_runtime("gemini")
        self.assertEqual(gemini.max_attempts, 1)
        self.assertEqual(gemini.models, ("model-utama",))

    async def test_starts_from_completed_audit(self):
        audit = SimpleNamespace(status="completed", source_ddl=ddl([1]), source_filename="schema.sql")
        request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(audit_semaphore=asyncio.Semaphore(1))))
        with patch.object(routes, "find_spreadsheet_job", return_value=audit), \
             patch.object(routes, "get_normalization_runtime", return_value=FakeRuntime("[]")), \
             patch.object(routes, "start_normalization_job", new_callable=AsyncMock, return_value="job") as start:
            created = await routes.create_normalization(routes.NormalizationRequest(audit_job_id="audit"), request)
        self.assertEqual(created["status_url"], "/api/v1/normalization/job")
        self.assertEqual(start.await_args.args[2], "audit")

    async def test_requires_existing_nonfailed_audit_before_start(self):
        data = routes.NormalizationRequest(audit_job_id="audit")
        with patch.object(routes, "find_spreadsheet_job", return_value=None):
            with self.assertRaises(HTTPException) as caught:
                await routes.create_normalization(data, None)
        self.assertEqual(caught.exception.status_code, 404)
        with patch.object(routes, "find_spreadsheet_job", return_value=SimpleNamespace(status="failed")):
            with self.assertRaises(HTTPException) as caught:
                await routes.create_normalization(data, None)
        self.assertEqual(caught.exception.status_code, 409)

    async def test_starts_while_audit_is_running(self):
        audit = SimpleNamespace(status="running", source_ddl=ddl([1]), source_filename="schema.sql",
                                normalization_job_id=None, normalization_error="sebelumnya gagal")
        request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(audit_semaphore=asyncio.Semaphore(1))))
        with patch.object(routes, "find_spreadsheet_job", return_value=audit), \
             patch.object(routes, "get_normalization_runtime", return_value=FakeRuntime("[]")), \
             patch.object(routes, "start_normalization_job", new_callable=AsyncMock, return_value="normal"):
            created = await routes.create_normalization(
                routes.NormalizationRequest(audit_job_id="audit"), request,
            )
        self.assertEqual(created["job_id"], "normal")
        self.assertEqual(audit.normalization_job_id, "normal")
        self.assertIsNone(audit.normalization_error)

    async def test_audit_upload_starts_both_jobs_and_exposes_normalization_status(self):
        audit = SimpleNamespace(normalization_job_id=None, normalization_error=None)
        upload = SimpleNamespace(filename="schema.sql", read=AsyncMock(return_value=ddl([1])))
        request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(audit_semaphore=asyncio.Semaphore(2))))
        with patch.object(audit_routes, "get_runtime", return_value=FakeRuntime("[]")), \
             patch.object(audit_routes, "get_normalization_runtime", return_value=FakeRuntime("[]")), \
             patch.object(audit_routes, "start_spreadsheet_job", new_callable=AsyncMock, return_value="audit"), \
             patch.object(audit_routes, "start_normalization_job", new_callable=AsyncMock, return_value="normal") as normal_start, \
             patch.object(audit_routes, "find_spreadsheet_job", return_value=audit):
            created = await audit_routes.create_audit_spreadsheet(request, upload, "telkom")
        self.assertEqual(created["normalization_status_url"], "/api/v1/normalization/normal")
        self.assertEqual(audit.normalization_job_id, "normal")
        self.assertEqual(normal_start.await_args.args[2], "audit")

    async def test_audit_and_normalization_run_concurrently(self):
        audit_started = asyncio.Event()
        normalization_started = asyncio.Event()
        release = asyncio.Event()

        async def run_audit(*_args, **_kwargs):
            audit_started.set()
            await release.wait()
            return b"audit-xlsx", {"findings": []}

        async def run_normalization(*_args, **_kwargs):
            normalization_started.set()
            await release.wait()
            return []

        upload = SimpleNamespace(filename="schema.sql", read=AsyncMock(return_value=ddl([1])))
        request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(audit_semaphore=asyncio.Semaphore(2))))
        with patch.dict(audit_pipeline._jobs, clear=True), patch.dict(pipeline._jobs, clear=True), \
             patch.object(audit_routes, "get_runtime", return_value=FakeRuntime("[]")), \
             patch.object(audit_routes, "get_normalization_runtime", return_value=FakeRuntime("[]")), \
             patch.object(audit_pipeline, "process_schema_audit", side_effect=run_audit), \
             patch.object(pipeline, "analyze_normalization", side_effect=run_normalization):
            created = await audit_routes.create_audit_spreadsheet(request, upload, "telkom")
            try:
                await asyncio.wait_for(
                    asyncio.gather(audit_started.wait(), normalization_started.wait()), timeout=2,
                )
                audit_status = await audit_routes.get_audit_spreadsheet_job(created["job_id"])
                self.assertEqual(audit_status["status"], "running")
                self.assertEqual(audit_status["normalization_status_url"], created["normalization_status_url"])
                self.assertEqual((await routes.get_normalization(created["normalization_job_id"]))["status"], "running")
                self.assertTrue(any(event["stage"] == "audit" for event in audit_status["process_events"]))
                self.assertTrue(any(event["stage"] == "normalization" for event in audit_status["process_events"]))
            finally:
                release.set()
                await asyncio.gather(*list(audit_pipeline._tasks), *list(pipeline._tasks))

    async def test_audit_remains_available_if_normalization_cannot_start(self):
        audit = SimpleNamespace(normalization_job_id=None, normalization_error=None)
        upload = SimpleNamespace(filename="schema.sql", read=AsyncMock(return_value=ddl([1])))
        request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(audit_semaphore=asyncio.Semaphore(2))))
        with patch.object(audit_routes, "get_runtime", return_value=FakeRuntime("[]")), \
             patch.object(audit_routes, "get_normalization_runtime", return_value=FakeRuntime("[]")), \
             patch.object(audit_routes, "start_spreadsheet_job", new_callable=AsyncMock, return_value="audit"), \
             patch.object(audit_routes, "start_normalization_job", new_callable=AsyncMock, side_effect=ValueError("gagal")), \
             patch.object(audit_routes, "find_spreadsheet_job", return_value=audit), \
             patch.object(audit_routes.logger, "exception"):
            created = await audit_routes.create_audit_spreadsheet(request, upload, "telkom")
        self.assertEqual(created["job_id"], "audit")
        self.assertIsNone(created["normalization_status_url"])
        self.assertIn("Coba lagi", created["normalization_error"])

    async def test_job_empty_result_download_headers_and_expiry(self):
        runtime = FakeRuntime("[]")
        source = ddl([1])
        with patch.dict(pipeline._jobs, clear=True):
            job_id = await pipeline.start_normalization_job(
                source, "schema.sql", "audit", runtime, asyncio.Semaphore(1),
            )
            await asyncio.gather(*list(pipeline._tasks))
            status = await routes.get_normalization(job_id)
            self.assertEqual(status["status"], "completed")
            self.assertEqual((status["completed_batches"], status["total_batches"]), (1, 1))
            self.assertEqual(len(runtime.calls), 1)
            self.assertEqual(status["finding_count"], 0)
            result = await routes.get_normalization_findings(job_id)
            self.assertEqual(result["findings"], [])
            response = await routes.download_normalization(job_id)
            self.assertIn("temuan_normalisasi.xlsx", response.headers["content-disposition"])
            sheet = load_workbook(BytesIO(response.body)).active
            self.assertEqual([cell.value for cell in sheet[1]], list(FINDING_COLUMNS))
            finding_sheet = load_workbook(BytesIO(create_normalization_excel([finding()]))).active
            self.assertEqual(finding_sheet["C2"].value, "pegawai")
            with patch.object(pipeline.time, "monotonic", return_value=pipeline._jobs[job_id].finished_at + 3601):
                self.assertIsNone(pipeline.find_normalization_job(job_id))

    async def test_dictionary_endpoints_require_matching_completed_normalization(self):
        audit = SimpleNamespace(status="completed", audit_result={"findings": []})
        old = SimpleNamespace(status="completed", result={}, audit_job_id="audit")
        new = SimpleNamespace(status="completed", result={}, audit_job_id="audit")
        classification = SimpleNamespace(status="completed", result={}, audit_job_id="audit")
        calls = [
            (dictionary_routes.create_classification,
             dictionary_routes.ClassificationRequest(audit_job_id="audit", normalization_job_id="normal")),
            (dictionary_routes.create_old_dictionary,
             dictionary_routes.OldDictionaryRequest(audit_job_id="audit", normalization_job_id="normal")),
            (dictionary_routes.create_new_dictionary,
             dictionary_routes.NewDictionaryRequest(audit_job_id="audit", normalization_job_id="normal",
                                                    classification_job_id="classification")),
            (dictionary_routes.create_log_perubahan,
             dictionary_routes.LogPerubahanRequest(audit_job_id="audit", normalization_job_id="normal",
                                                   old_dictionary_job_id="old", new_dictionary_job_id="new")),
        ]
        with patch.object(dictionary_routes, "find_spreadsheet_job", return_value=audit), \
             patch.object(dictionary_routes, "find_classification_job", return_value=classification), \
             patch.object(dictionary_routes, "find_old_dictionary_job", return_value=old), \
             patch.object(dictionary_routes, "find_new_dictionary_job", return_value=new):
            for job, expected_code in ((None, 404),
                                       (SimpleNamespace(status="running", audit_job_id="audit"), 409),
                                       (SimpleNamespace(status="completed", audit_job_id="other"), 422)):
                with patch.object(dictionary_routes, "find_normalization_job", return_value=job):
                    for endpoint, data in calls:
                        with self.subTest(endpoint=endpoint.__name__, code=expected_code):
                            with self.assertRaises(HTTPException) as caught:
                                await endpoint(data, None)
                            self.assertEqual(caught.exception.status_code, expected_code)

    async def test_dictionary_new_receives_completed_normalization_findings(self):
        source = ddl([1])
        audit = SimpleNamespace(status="completed", source_ddl=source,
                                audit_result={"findings": [{"nama_tabel": "table_0"}]})
        normalization = SimpleNamespace(status="completed", audit_job_id="audit",
                                        findings=[finding("table_0", "field_0")])
        classification = SimpleNamespace(status="completed", audit_job_id="audit", result={"table_0": {}})
        request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(audit_semaphore=asyncio.Semaphore(1))))
        data = dictionary_routes.NewDictionaryRequest(
            audit_job_id="audit", normalization_job_id="normal", classification_job_id="classification",
            model="gpt_setneg",
        )
        with patch.object(dictionary_routes, "find_spreadsheet_job", return_value=audit), \
             patch.object(dictionary_routes, "find_normalization_job", return_value=normalization), \
             patch.object(dictionary_routes, "find_classification_job", return_value=classification), \
             patch.object(dictionary_routes, "start_new_dictionary_job", new_callable=AsyncMock,
                          return_value="new-job") as start:
            response = await dictionary_routes.create_new_dictionary(data, request)
        self.assertEqual(response["job_id"], "new-job")
        self.assertIs(start.await_args.args[2], normalization.findings)
        self.assertIs(start.await_args.args[3], classification.result)

        normalization.findings = []
        with patch.object(dictionary_routes, "find_spreadsheet_job", return_value=audit), \
             patch.object(dictionary_routes, "find_normalization_job", return_value=normalization), \
             patch.object(dictionary_routes, "find_classification_job", return_value=classification), \
             patch.object(dictionary_routes, "start_new_dictionary_job", new_callable=AsyncMock,
                          return_value="new-job") as start:
            await dictionary_routes.create_new_dictionary(data, request)
        self.assertEqual(start.await_args.args[2], [])

        normalization.findings = None
        with patch.object(dictionary_routes, "find_spreadsheet_job", return_value=audit), \
             patch.object(dictionary_routes, "find_normalization_job", return_value=normalization), \
             patch.object(dictionary_routes, "find_classification_job", return_value=classification):
            with self.assertRaises(HTTPException) as caught:
                await dictionary_routes.create_new_dictionary(data, request)
        self.assertEqual(caught.exception.status_code, 409)


if __name__ == "__main__":
    unittest.main()
