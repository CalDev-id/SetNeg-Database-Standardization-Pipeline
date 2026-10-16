import asyncio
import json
import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from features.audit_db.pipeline import process_schema_audit
from features.audit_db.cari_temuan import DISCOVERY_SYSTEM_PROMPT, baca_kandidat, buat_batch_audit
from features.audit_db.buat_solusi import SOLUTION_SYSTEM_PROMPT, baca_narasi
from features.audit_db.verifikasi_temuan import VERIFICATION_SYSTEM_PROMPT
from playbook.playbook import VALID_OBJECT_NAMES, format_playbook_rule
from shared.ddl.parser import parse_schema_ddl
from shared.json_utils import JsonFormatError, baca_json, to_json
from shared.llm.groq_runtime import GroqPayloadTooLargeError, GroqRuntime
from shared.llm.runtime import PayloadTooLargeError


DDL = b"""
CREATE TABLE app_dev.person (
  id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
  email VARCHAR(255),
  PRIMARY KEY (id)
);
"""
MANY_DDL = (
    b"CREATE TABLE app_dev.sample (id BIGINT PRIMARY KEY, "
    + b", ".join(f"`BAD_{index}` VARCHAR(10)".encode() for index in range(13))
    + b");"
)


class FakeRuntime:
    batch_size = 50

    def __init__(self, no_findings=False):
        self.calls = []
        self.no_findings = no_findings

    def generate_response(self, system_prompt, user_prompt):
        if "TAHAP 1: CARI TEMUAN" in system_prompt:
            self.calls.append("find")
            if self.no_findings:
                return "[]"
            candidate = {
                "Nama Tabel": "person", "Objek Temuan": "Nama Primary Key",
                "Aturan Playbook/UU": "2b. Penamaan objek spesifik",
                "Objek SQL": "id", "Bukti DDL": "PRIMARY KEY (id)",
                "Alasan": "Nama PK harus person_id.",
            }
            other = {
                "Nama Tabel": "person", "Objek Temuan": "COMMENT tabel",
                "Aturan Playbook/UU": "3b. Dokumentasi", "Objek SQL": "person",
                "Bukti DDL": "Opsi CREATE TABLE tidak memiliki COMMENT.",
                "Alasan": "COMMENT tabel wajib ada.",
            }
            return json.dumps([candidate, candidate, other])
        if "TAHAP 2: VERIFIKASI" in system_prompt:
            self.calls.append("verify")
            assert '"id": 1' in user_prompt and '"id": 2' in user_prompt
            return json.dumps({"decisions": [
                {"id": 1, "valid": True, "status_pii": "Tidak"},
                {"id": 2, "valid": False, "status_pii": "Tidak"},
            ]})
        self.calls.append("solve")
        assert "TAHAP 3: TULIS HASIL" in system_prompt
        return json.dumps([{
            "id": 1, "Nama Tabel": "invented_table",
            "Uraian Ketidaksesuaian": "PK `id` pada `person` tidak sesuai nama baku.",
            "Solusi": "ALTER TABLE `person` RENAME COLUMN `id` TO `person_id`;",
            "Catatan": "-",
        }])


class SplitRuntime(FakeRuntime):
    def generate_response(self, system_prompt, user_prompt):
        if "TAHAP 1: CARI TEMUAN" in system_prompt:
            self.calls.append("find")
            batch_json = user_prompt.split("Data batch:\n", 1)[1]
            batch, _ = json.JSONDecoder().raw_decode(batch_json)
            columns = sum(len(table["definisi_kolom_batch"]) for table in batch["tabel"])
            if columns > 1:
                return '[{"Nama Tabel":"person"}'
            if not batch["tabel"][0]["batch_pertama"]:
                return "[]"
            return json.dumps([{
                "Nama Tabel": "person", "Objek Temuan": "Nama Primary Key",
                "Aturan Playbook/UU": "2b. Penamaan objek spesifik",
                "Objek SQL": "id", "Bukti DDL": "PRIMARY KEY (id)",
                "Alasan": "Nama PK harus person_id.",
            }])
        if "TAHAP 2: VERIFIKASI" in system_prompt:
            self.calls.append("verify")
            return json.dumps({"decisions": [{"id": 1, "valid": True, "status_pii": "Tidak"}]})
        self.calls.append("solve")
        return json.dumps([{
            "id": 1,
            "Uraian Ketidaksesuaian": "PK `id` pada `person` tidak sesuai nama baku.",
            "Solusi": "ALTER TABLE `person` RENAME COLUMN `id` TO `person_id`;", "Catatan": "-",
        }])


class PayloadSplitRuntime(FakeRuntime):
    def generate_response(self, system_prompt, user_prompt):
        if "TAHAP 1: CARI TEMUAN" in system_prompt:
            self.calls.append("find")
            batch_json = user_prompt.split("Data batch:\n", 1)[1]
            batch, _ = json.JSONDecoder().raw_decode(batch_json)
            columns = sum(len(table["definisi_kolom_batch"]) for table in batch["tabel"])
            if columns > 1:
                raise PayloadTooLargeError("Payload terlalu besar pada tahap pencarian temuan.")
            self.calls.pop()
        return super().generate_response(system_prompt, user_prompt)


class ManyRuntime:
    batch_size = 50

    def __init__(self):
        self.calls = []

    def generate_response(self, system_prompt, user_prompt):
        if "TAHAP 1: CARI TEMUAN" in system_prompt:
            self.calls.append("find")
            return json.dumps([{
                "Nama Tabel": "sample", "Objek Temuan": "Nama kolom",
                "Aturan Playbook/UU": "2a. Penamaan universal", "Objek SQL": f"BAD_{index}",
                "Bukti DDL": f"`BAD_{index}` VARCHAR(10)", "Alasan": "Nama harus lowercase.",
            } for index in range(13)])
        if "TAHAP 2: VERIFIKASI" in system_prompt:
            self.calls.append("verify")
            items = json.loads(user_prompt.split("Kandidat yang harus diverifikasi:\n", 1)[1])
            return json.dumps({"decisions": [
                {"id": item["id"], "valid": True, "status_pii": "Tidak"} for item in items
            ]})
        self.calls.append("solve")
        items = json.loads(user_prompt.split("Kandidat yang sudah lolos verifikasi:\n", 1)[1])
        return json.dumps([{
            "id": item["id"], "Uraian Ketidaksesuaian": f"Kolom {item['Objek SQL']} tidak lowercase.",
            "Solusi": f"Ganti nama kolom {item['Objek SQL']} menjadi lowercase.", "Catatan": "-",
        } for item in items])


class AuditPipelineTest(unittest.TestCase):
    def test_shared_json_checker_and_serializer(self):
        payload = {"nama": "uji", "aktif": True}
        self.assertEqual(baca_json(to_json(payload)), payload)
        self.assertEqual(baca_json("```json\n[1, 2]\n```"), [1, 2])

    def test_playbook_rule_is_always_displayed_with_its_title(self):
        self.assertEqual(format_playbook_rule("3b"), "3b. Dokumentasi")
        self.assertEqual(
            format_playbook_rule("2b Penamaan objek spesifik"),
            "2b. Penamaan objek spesifik",
        )
        self.assertIsNone(format_playbook_rule("5. Aturan lain"))

    def test_primary_key_rename_uses_direct_column_rename(self):
        valid = json.dumps([{
            "id": 1,
            "Uraian Ketidaksesuaian": "Nama PK tidak sesuai.",
            "Solusi": "ALTER TABLE `person` RENAME COLUMN `id` TO `person_id`;",
            "Catatan": "-",
        }])
        self.assertIn(1, baca_narasi(valid, {1}, {1}))

        invalid = valid.replace(
            "ALTER TABLE `person` RENAME COLUMN `id` TO `person_id`;",
            "Contoh SQL: ALTER TABLE person DROP PRIMARY KEY, ADD PRIMARY KEY (person_id);",
        )
        with self.assertRaises(JsonFormatError):
            baca_narasi(invalid, {1}, {1})

    def test_groq_uses_qwen_when_primary_model_is_limited(self):
        class Response:
            def __init__(self, status_code, payload=None, headers=None):
                self.status_code = status_code
                self.payload = payload
                self.headers = headers or {}

            def json(self):
                return self.payload

        with patch.dict("os.environ", {"GROQ_API_KEY": "test-key"}):
            with patch(
                "shared.llm.groq_runtime.requests.post",
                side_effect=[
                    Response(429),
                    Response(200, {"choices": [{"message": {"content": "[]"}}]}),
                ],
            ) as post:
                runtime = GroqRuntime()
                self.assertEqual(runtime.generate_response("system", "user"), "[]")

        self.assertEqual(runtime.model, "qwen/qwen3.8-27b")
        self.assertEqual(
            [call.kwargs["json"]["model"] for call in post.call_args_list],
            ["openai/gpt-oss-120b", "qwen/qwen3.8-27b"],
        )

    def test_groq_retries_after_both_models_are_limited(self):
        class Response:
            def __init__(self, status_code, payload=None):
                self.status_code = status_code
                self.payload = payload
                self.headers = {"Retry-After": "1"}

            def json(self):
                return self.payload

        with patch.dict("os.environ", {"GROQ_API_KEY": "test-key"}):
            with patch(
                "shared.llm.groq_runtime.requests.post",
                side_effect=[
                    Response(429),
                    Response(429),
                    Response(200, {"choices": [{"message": {"content": "[]"}}]}),
                ],
            ) as post, patch("shared.llm.groq_runtime.time.sleep") as sleep:
                runtime = GroqRuntime()
                self.assertEqual(runtime.generate_response("system", "user"), "[]")

        self.assertEqual(sleep.call_args.args, (1,))
        self.assertEqual(
            [call.kwargs["json"]["model"] for call in post.call_args_list],
            ["openai/gpt-oss-120b", "qwen/qwen3.8-27b", "openai/gpt-oss-120b"],
        )

    def test_groq_marks_http_413_as_payload_too_large(self):
        class Response:
            status_code = 413
            headers = {}

        with patch.dict("os.environ", {"GROQ_API_KEY": "test-key"}):
            with patch(
                "shared.llm.groq_runtime.requests.post", return_value=Response(),
            ):
                with self.assertRaises(GroqPayloadTooLargeError):
                    GroqRuntime().generate_response("system", "user")

    def test_prompts_use_original_playbook_rules(self):
        for prompt in (DISCOVERY_SYSTEM_PROMPT, VERIFICATION_SYSTEM_PROMPT):
            for number in range(1, 13):
                self.assertIn(f"C{number}.", prompt)
            self.assertIn("D. KONSISTENSI DAN LARANGAN", prompt)
            self.assertIn("F. NOMOR ATURAN OUTPUT", prompt)
            self.assertNotIn("H. FORMAT OUTPUT", prompt)
            self.assertIn("Nama Primary Key, nama tabel/database/FK/index/unique/check/view memakai 2b", prompt)
            self.assertIn("Tipe Primary Key dan tipe/ukuran data memakai 2c", prompt)
        self.assertNotIn("C12. KOLOM PII", SOLUTION_SYSTEM_PROMPT)
        self.assertIn("TAHAP 3: TULIS HASIL", SOLUTION_SYSTEM_PROMPT)
        self.assertIn("Nama Primary Key", VALID_OBJECT_NAMES)

    def test_three_stages_lock_findings_and_keep_excel_columns(self):
        runtime = FakeRuntime()
        progress = []
        result = asyncio.run(process_schema_audit(
            DDL, "json", runtime=runtime, verification_runtime=runtime,
            on_batch_complete=lambda: progress.append(1),
        ))
        self.assertEqual(runtime.calls, ["find", "verify", "solve"])
        self.assertEqual(len(progress), 3)
        self.assertEqual(len(result["findings"]), 1)
        self.assertEqual(result["findings"][0]["nama_tabel"], "person")
        self.assertEqual(result["findings"][0]["objek_temuan"], "Nama Primary Key")
        self.assertEqual(result["findings"][0]["no"], 1)

        excel_runtime = FakeRuntime()
        excel = asyncio.run(process_schema_audit(
            DDL, "excel", runtime=excel_runtime, verification_runtime=excel_runtime,
        ))
        sheet = load_workbook(BytesIO(excel))["MySQL - Setneg Point"]
        self.assertEqual([sheet.cell(8, column).value for column in range(2, 10)], [
            "No", "Nama Database", "Nama Tabel", "Objek Temuan", "Aturan Playbook/UU",
            "Uraian Ketidaksesuaian", "Solusi", "Catatan",
        ])
        self.assertEqual(sheet["D9"].value, "person")
        self.assertEqual(sheet["E9"].value, "Nama Primary Key")

    def test_empty_findings_skip_later_requests_and_incomplete_json_is_rejected(self):
        runtime = FakeRuntime(no_findings=True)
        progress = []
        result = asyncio.run(process_schema_audit(
            DDL, "json", runtime=runtime, verification_runtime=runtime,
            on_batch_complete=lambda: progress.append(1),
        ))
        self.assertEqual(result["findings"], [])
        self.assertEqual(runtime.calls, ["find"])
        self.assertEqual(len(progress), 3)
        parsed = parse_schema_ddl(DDL)
        parts = buat_batch_audit(parsed, 50)[0]
        with self.assertRaises(JsonFormatError):
            baca_kandidat('[{"Nama Tabel":"person"}', parts, True)
        with self.assertRaises(JsonFormatError):
            baca_narasi("[]", {1})

    def test_incomplete_discovery_retries_with_smaller_column_batches(self):
        runtime = SplitRuntime()
        result = asyncio.run(process_schema_audit(
            DDL, "json", runtime=runtime, verification_runtime=runtime,
        ))
        self.assertEqual(runtime.calls, ["find", "find", "find", "find", "verify", "solve"])
        self.assertEqual(len(result["findings"]), 1)
        self.assertEqual(result["findings"][0]["objek_temuan"], "Nama Primary Key")

    def test_payload_too_large_discovery_is_split_before_failing(self):
        runtime = PayloadSplitRuntime()
        result = asyncio.run(process_schema_audit(
            DDL, "json", runtime=runtime, verification_runtime=runtime,
        ))
        self.assertEqual(runtime.calls, ["find", "find", "find", "verify", "solve"])
        self.assertEqual(len(result["findings"]), 1)

    def test_long_solution_stage_is_split_without_losing_findings(self):
        runtime = ManyRuntime()
        result = asyncio.run(process_schema_audit(
            MANY_DDL, "json", runtime=runtime, verification_runtime=runtime,
        ))
        self.assertEqual(runtime.calls, ["find", "verify", "solve", "solve"])
        self.assertEqual(len(result["findings"]), 13)
        self.assertEqual([item["no"] for item in result["findings"]], list(range(1, 14)))


if __name__ == "__main__":
    unittest.main()
