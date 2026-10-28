"""Ekspor temuan normalisasi menjadi Excel."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from features.normalization.analysis import FINDING_COLUMNS


def create_normalization_excel(findings: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Temuan Normalisasi"
    sheet.append(FINDING_COLUMNS)
    for finding in findings:
        sheet.append([finding[column] for column in FINDING_COLUMNS])
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"A1:K{max(1, sheet.max_row)}"
    widths = [7, 22, 24, 19, 34, 28, 55, 55, 38, 13, 35]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="173B65")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
