import asyncio
from copy import copy
from io import BytesIO
from pathlib import Path
from textwrap import wrap
from typing import Callable, Optional, Union

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from features.audit_db.buat_solusi import buat_solusi
from features.audit_db.cari_temuan import (
    BATCH_SIZE, build_user_prompt, buat_batch_audit,
    candidate_key, cari_temuan,
)
from features.audit_db.verifikasi_temuan import verifikasi_temuan
from shared.ddl.parser import ParsedDdl, parse_schema_ddl
from shared.llm.gemini_runtime import GeminiRuntime
from shared.llm.groq_runtime import GroqRuntime
from shared.llm.runtime import AuditError, LlmRuntime


AUDIT_STAGES_PER_BATCH = 3


class Temuan(BaseModel):
    no: int = Field(ge=1, alias="No")
    database_name: str = Field(alias="Nama Database")
    table_name: str = Field(alias="Nama Tabel")
    object_name: str = Field(alias="Objek Temuan")
    playbook_rule: str = Field(alias="Aturan Playbook/UU")
    nonconformity: str = Field(alias="Uraian Ketidaksesuaian")
    solution: str = Field(alias="Solusi")
    status: str = Field(alias="Status")
    note: str = Field(alias="Catatan")
    pii_status: str = Field(alias="Status PII")


async def audit_ddl(
    parsed_ddl: ParsedDdl, runtime: Optional[LlmRuntime] = None,
    on_batch_complete: Optional[Callable[[], None]] = None,
    verification_runtime: Optional[LlmRuntime] = None,
) -> list[Temuan]:
    if runtime is None:
        try:
            runtime = GroqRuntime()
        except ValueError as error:
            raise AuditError(str(error)) from error
    batch_size = max(1, int(getattr(runtime, "batch_size", BATCH_SIZE)))
    all_table_names = list(parsed_ddl.tables)
    findings: list[Temuan] = []
    seen_accepted: set[tuple[str, str, str]] = set()
    batches = buat_batch_audit(parsed_ddl, batch_size, getattr(runtime, "max_tables_per_batch", None))
    for batch_index, batch in enumerate(batches):
        check_database = batch_index == 0 and bool(parsed_ddl.database_name)
        batch_prompt = build_user_prompt(batch, parsed_ddl.database_name, check_database, all_table_names)
        candidates = await cari_temuan(parsed_ddl, runtime, batch, check_database, all_table_names)
        if on_batch_complete:
            on_batch_complete()
        numbered = list(enumerate(candidates, start=1))
        if numbered and verification_runtime is None:
            try:
                verification_runtime = GeminiRuntime(model="gemini-3.5-flash-lite")
            except ValueError as error:
                raise AuditError(
                    "GEMINI_API_KEY diperlukan untuk tahap verifikasi temuan: "
                    f"{error}"
                ) from error
        decisions = (
            await verifikasi_temuan(verification_runtime, numbered, batch_prompt)
            if numbered else {}
        )
        if on_batch_complete:
            on_batch_complete()
        accepted = []
        batch_keys = set()
        for identifier, candidate in numbered:
            key = candidate_key(candidate)
            if decisions[identifier].valid and key not in seen_accepted and key not in batch_keys:
                accepted.append((identifier, candidate))
                batch_keys.add(key)
        narratives = await buat_solusi(runtime, accepted) if accepted else {}
        if on_batch_complete:
            on_batch_complete()
        for identifier, candidate in accepted:
            narrative = narratives[identifier]
            seen_accepted.add(candidate_key(candidate))
            findings.append(Temuan.model_validate({
                "No": 1,
                "Nama Database": parsed_ddl.database_name or "-",
                "Nama Tabel": candidate.table_name,
                "Objek Temuan": candidate.object_name,
                "Aturan Playbook/UU": candidate.playbook_rule,
                "Uraian Ketidaksesuaian": narrative.nonconformity,
                "Solusi": narrative.solution,
                "Status": "To Do",
                "Catatan": narrative.note or "-",
                "Status PII": decisions[identifier].status_pii,
            }))

    table_order = {name: index for index, name in enumerate(parsed_ddl.tables)}
    findings.sort(key=lambda item: -1 if item.table_name == "-" else table_order[item.table_name])
    return findings


TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "template" / "template.xlsx"
FIRST_DATA_ROW = 9


def buat_respons_audit(parsed_ddl: ParsedDdl, temuan: list[Temuan]) -> dict:
    return {
        "message": "Analisis standardisasi selesai.",
        "summary": {
            "table_count": len(parsed_ddl.tables),
            "column_count": parsed_ddl.column_count,
            "tables": sorted(parsed_ddl.tables),
        },
        "findings": [
            {
                "no": index,
                "nama_tabel": item.table_name,
                "objek_temuan": item.object_name,
                "aturan_playbook": item.playbook_rule,
                "uraian_ketidaksesuaian": item.nonconformity,
                "solusi": item.solution,
                "status": item.status,
                "catatan": item.note,
                "status_pii": item.pii_status,
            }
            for index, item in enumerate(temuan, start=1)
        ],
    }


def _cell(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def buat_excel_temuan(database_name: str, temuan: list[Temuan]) -> bytes:
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook["MySQL - Setneg Point"]
    sheet.column_dimensions["H"].width = 55
    sheet.column_dimensions["I"].hidden = False
    sheet.column_dimensions["I"].width = sheet.column_dimensions["J"].width
    sheet.column_dimensions["J"].hidden = True
    sheet["I8"] = "Catatan"
    sheet["I8"]._style = copy(sheet["J8"]._style)
    sheet["J8"] = None

    sheet.data_validations.dataValidation.clear()

    last_row = FIRST_DATA_ROW + len(temuan) - 1
    for index, item in enumerate(temuan, start=1):
        row_number = FIRST_DATA_ROW + index - 1
        values = (
            index, _cell(database_name), _cell(item.table_name),
            _cell(item.object_name), _cell(item.playbook_rule),
            _cell(item.nonconformity), _cell(item.solution), _cell(item.note),
        )
        max_lines = 1
        for column, value in enumerate(values, start=2):
            source_column = column if column < 9 else 10
            cell = sheet.cell(row_number, column, value)
            cell._style = copy(sheet.cell(FIRST_DATA_ROW, source_column)._style)
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            cell.alignment = alignment
            width = max(1, int(sheet.column_dimensions[get_column_letter(column)].width) - 2)
            lines = sum(max(1, len(wrap(part, width=width))) for part in str(value).split("\n"))
            max_lines = max(max_lines, lines)
        sheet.row_dimensions[row_number].height = min(409.5, max(15.75, max_lines * 15))
        sheet.cell(row_number, 10).value = None

    rules = list(sheet.conditional_formatting["B9:J42"])
    del sheet.conditional_formatting["B9:J42"]
    for rule in rules:
        sheet.conditional_formatting.add(f"B9:I{max(last_row, FIRST_DATA_ROW)}", rule)

    total = len(temuan)
    selesai = sum(item.status == "Done" for item in temuan)
    sheet["B5"] = "TOTAL TEMUAN"
    sheet["E5"] = "TEMUAN AKTIF"
    sheet["G5"] = "TEMUAN SELESAI"
    sheet["I5"] = "TINGKAT PENYELESAIAN"
    sheet["B6"] = total
    sheet["E6"] = total - selesai
    sheet["G6"] = selesai
    sheet["I6"] = selesai / total if total else 0
    sheet["I6"].number_format = "0.0%"
    sheet.auto_filter.ref = f"B8:I{max(last_row, FIRST_DATA_ROW)}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def process_schema_audit(
    sql_input: Union[bytes, ParsedDdl], output: str,
    on_batch_complete: Optional[Callable[[], None]] = None,
    runtime: Optional[LlmRuntime] = None,
    verification_runtime: Optional[LlmRuntime] = None,
):
    parsed_ddl = (
        sql_input if isinstance(sql_input, ParsedDdl)
        else await asyncio.to_thread(parse_schema_ddl, sql_input)
    )
    temuan = await audit_ddl(
        parsed_ddl,
        runtime=runtime,
        on_batch_complete=on_batch_complete,
        verification_runtime=verification_runtime,
    )
    if output == "excel":
        return await asyncio.to_thread(buat_excel_temuan, parsed_ddl.database_name or "-", temuan)
    if output == "json":
        return buat_respons_audit(parsed_ddl, temuan)
    raise ValueError("Format output tidak dikenal.")
