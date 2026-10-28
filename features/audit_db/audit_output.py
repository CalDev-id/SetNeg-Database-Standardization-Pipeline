"""Pembentukan respons dan spreadsheet hasil audit database."""

from io import BytesIO
from copy import copy
from pathlib import Path
from textwrap import wrap
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from features.audit_db.audit import Temuan
from features.audit_db.extract import ParsedDdl


TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "template" / "template.xlsx"
FIRST_DATA_ROW = 9
COMBINED_COLUMNS = (
    ("No", "no"), ("Nama Database", "nama_database"),
    ("Nama Tabel", "nama_tabel"), ("Objek Temuan", "objek_temuan"),
    ("Aturan Playbook/UU", "aturan_playbook"),
    ("Uraian Ketidaksesuaian", "uraian_ketidaksesuaian"),
    ("Solusi", "solusi"), ("Status", "status"),
    ("Catatan", "catatan"), ("Jenis Temuan", "jenis_temuan"),
)
COMBINED_ROW_FILLS = {
    "normalization": PatternFill("solid", fgColor="FFEAF4FF"),
    "removal": PatternFill("solid", fgColor="FFFFE2DF"),
    "addition": PatternFill("solid", fgColor="FFFFF1BF"),
}


def buat_respons_audit(parsed_ddl: ParsedDdl, temuan: list[Temuan]) -> dict:
    return {
        "message": "Analisis standardisasi selesai.",
        "summary": {
            "table_count": len(parsed_ddl.tables),
            "column_count": parsed_ddl.column_count,
            "tables": sorted(parsed_ddl.tables),
        },
        "findings": [
            {
                "no": index,
                "nama_database": item.database_name,
                "nama_tabel": item.table_name,
                "objek_temuan": item.object_name,
                "aturan_playbook": item.playbook_rule,
                "uraian_ketidaksesuaian": item.nonconformity,
                "solusi": item.solution,
                "status": item.status,
                "catatan": item.note,
            }
            for index, item in enumerate(temuan, start=1)
        ],
    }


def _cell(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def combined_findings(audit_findings: list[dict], normalization_findings: list[dict]) -> list[dict]:
    rows = []
    for finding in audit_findings:
        rows.append({**finding, "no": len(rows) + 1, "jenis_temuan": "Audit Playbook"})
    for finding in normalization_findings:
        solution = finding["Solusi Normalisasi"]
        proposed = finding["Tabel Baru Usulan"].strip()
        if proposed and proposed != "-":
            solution += f"\nTabel baru usulan: {proposed}"
        rows.append({
            "no": len(rows) + 1,
            "nama_database": finding["Nama Database"],
            "nama_tabel": finding["Nama Tabel"],
            "objek_temuan": finding["Objek Temuan"],
            "aturan_playbook": f"Normalisasi Database ({finding['Tahap Normalisasi']})",
            "uraian_ketidaksesuaian": (
                f"Kolom terdampak: {finding['Kolom Terdampak']}\n"
                f"{finding['Uraian Ketidaksesuaian']}"
            ),
            "solusi": solution,
            "status": finding["Status"],
            "catatan": finding["Catatan"],
            "jenis_temuan": "Normalisasi",
        })
    return rows


def _combined_row_fill(finding: dict) -> Optional[PatternFill]:
    if finding["jenis_temuan"] == "Normalisasi":
        return COMBINED_ROW_FILLS["normalization"]
    if finding["jenis_temuan"] != "Audit Playbook":
        return None

    object_name = str(finding["objek_temuan"]).lower()
    detail = str(finding["uraian_ketidaksesuaian"]).lower()
    solution = str(finding["solusi"]).lower()
    combined = object_name + detail + solution
    if "drop" in combined or "hapus" in combined:
        return COMBINED_ROW_FILLS["removal"]
    if any(word in object_name for word in ("comment", "foreign key", "index", "check")):
        return None
    adds_object = any(word in object_name for word in ("tambah", "penambahan")) and any(
        word in object_name for word in ("kolom", "tabel")
    )
    adds_from_solution = any(phrase in solution for phrase in (
        "tambahkan kolom", "tambahkan tabel", "tambahkan satu tabel", "tambahkan satu set audit trail",
    ))
    adds_from_description = any(phrase in detail + solution for phrase in (
        "penambahan kolom", "penambahan tabel",
    ))
    if ("audit trail" in object_name or "tabel metadata" in object_name or
            adds_object or adds_from_solution or adds_from_description):
        return COMBINED_ROW_FILLS["addition"]
    return None


def buat_excel_gabungan(findings: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Temuan Database"
    sheet.append([label for label, _ in COMBINED_COLUMNS])
    widths = (8, 24, 25, 38, 34, 60, 60, 14, 42, 18)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="173B65")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for finding in findings:
        sheet.append([
            finding[key] if isinstance(finding[key], int) else _cell(str(finding[key]))
            for _, key in COMBINED_COLUMNS
        ])
        row = sheet[sheet.max_row]
        fill = _combined_row_fill(finding)
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill is not None:
                cell.fill = fill
        lines = max(
            sum(max(1, len(wrap(part, width=width))) for part in str(cell.value).split("\n"))
            for cell, width in zip(row, widths)
        )
        sheet.row_dimensions[row[0].row].height = min(409.5, max(18, lines * 15))
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"A1:J{max(1, sheet.max_row)}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def buat_excel_temuan(database_name: str, temuan: list[Temuan]) -> bytes:
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook["MySQL - Setneg Point"]
    sheet.column_dimensions["H"].width = 55
    sheet.column_dimensions["I"].hidden = False
    sheet.column_dimensions["I"].width = sheet.column_dimensions["J"].width
    sheet.column_dimensions["J"].hidden = True
    sheet["I8"] = "Catatan"
    sheet["I8"]._style = copy(sheet["J8"]._style)
    sheet["J8"] = None

    sheet.data_validations.dataValidation.clear()

    last_row = FIRST_DATA_ROW + len(temuan) - 1
    for index, item in enumerate(temuan, start=1):
        row_number = FIRST_DATA_ROW + index - 1
        values = (
            index, _cell(database_name), _cell(item.table_name),
            _cell(item.object_name), _cell(item.playbook_rule),
            _cell(item.nonconformity), _cell(item.solution), _cell(item.note),
        )
        max_lines = 1
        for column, value in enumerate(values, start=2):
            source_column = column if column < 9 else 10
            cell = sheet.cell(row_number, column, value)
            cell._style = copy(sheet.cell(FIRST_DATA_ROW, source_column)._style)
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            cell.alignment = alignment
            width = max(1, int(sheet.column_dimensions[get_column_letter(column)].width) - 2)
            lines = sum(max(1, len(wrap(part, width=width))) for part in str(value).split("\n"))
            max_lines = max(max_lines, lines)
        sheet.row_dimensions[row_number].height = min(409.5, max(15.75, max_lines * 15))
        sheet.cell(row_number, 10).value = None

    rules = list(sheet.conditional_formatting["B9:J42"])
    del sheet.conditional_formatting["B9:J42"]
    for rule in rules:
        sheet.conditional_formatting.add(f"B9:I{max(last_row, FIRST_DATA_ROW)}", rule)

    total = len(temuan)
    selesai = sum(item.status == "Done" for item in temuan)
    sheet["B5"] = "TOTAL TEMUAN"
    sheet["E5"] = "TEMUAN AKTIF"
    sheet["G5"] = "TEMUAN SELESAI"
    sheet["I5"] = "TINGKAT PENYELESAIAN"
    sheet["B6"] = total
    sheet["E6"] = total - selesai
    sheet["G6"] = selesai
    sheet["I6"] = selesai / total if total else 0
    sheet["I6"].number_format = "0.0%"
    sheet.auto_filter.ref = f"B8:I{max(last_row, FIRST_DATA_ROW)}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
